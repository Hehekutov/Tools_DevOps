"""Build academic reports from verified evidence and a structured content file."""
from pathlib import Path
import json,re,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageChops

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/reports';OUT.mkdir(parents=True,exist_ok=True)
EVID=ROOT/'reports/practices-5-8/evidence'
QUESTIONS=json.loads((ROOT/'lab_tools/questions.json').read_text(encoding='utf-8'))

def para(doc,text='',style=None):
    p=doc.add_paragraph(text,style);p.paragraph_format.widow_control=True
    return p

def code(doc,text):
    lines=text.rstrip().splitlines()
    for i,line in enumerate(lines):
        # Soft wrapping keeps terminal output inside the printable page.
        p=doc.add_paragraph(style='Code');p.add_run(line.expandtabs(4))
        if len(lines)<=24:p.paragraph_format.keep_with_next=i<len(lines)-1

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for cell,value in zip(t.rows[0].cells,headers):cell.text=value
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                for run in p.runs:run.font.size=Pt(10);run.bold=ri==0
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    para(doc)

def create(n,spec):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.5);s.right_margin=Cm(1.5)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
        props=st.element.find(qn('w:pPr'))
        if props is not None:
            for border in list(props.findall(qn('w:pBdr'))):props.remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for name,size in [('Title',18),('Heading 1',15),('Heading 2',13)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    st=d.styles.add_style('Code',1);st.font.name='Consolas';st.font.size=Pt(8)
    st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.keep_together=True
    s.different_first_page_header_footer=True
    cover_footer=s.first_page_footer.paragraphs[0];cover_footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=cover_footer.add_run('2026');run.font.name='Times New Roman';run.font.size=Pt(12)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for text in ['ИНСТИТУТ ИПТИП','Кафедра индустриального программирования','Дисциплина Инструменты девопс']:
        p=para(d,text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(3):para(d)
    p=para(d,f'Отчёт по практической работе {n}','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=para(d,spec['title'],'Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(4):para(d)
    for text in ['Выполнил Кудояров Денис Рустемович','Группа ЭФБО-18-24','Проверяющий Гиматдинов Дамир Маратович']:
        p=para(d,text);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    d.add_page_break()
    d.add_heading('Цель и задачи работы',level=1)
    for text in spec['intro']:para(d,text)
    d.add_heading('Рабочее окружение',level=1)
    table(d,['Параметр','Значение'],spec['environment'],[5,12])
    for section in spec['sections']:
        d.add_heading(section['title'],level=1)
        for text in section.get('paragraphs',[]):para(d,text)
        if 'table' in section:
            obj=section['table'];table(d,obj['headers'],obj['rows'],obj.get('widths'))
        if 'commands' in section:code(d,section['commands'])
        for item in section.get('evidence',[]):
            p=EVID/item['file']
            if not p.exists():raise FileNotFoundError(p)
            para(d,item['caption']).paragraph_format.keep_with_next=True
            text=p.read_text(encoding='utf-8',errors='replace')
            if 'start' in item:
                if item['start'] not in text:raise ValueError('Missing evidence marker: '+item['start'])
                text=text.split(item['start'],1)[1].lstrip('\n')
            if 'end' in item:
                if item['end'] not in text:raise ValueError('Missing evidence marker: '+item['end'])
                text=text.split(item['end'],1)[0].rstrip()
            if 'lines' in item:text='\n'.join(text.splitlines()[:item['lines']])
            code(d,text)
        for item in section.get('images',[]):
            p=EVID/item['file']
            if not p.exists():raise FileNotFoundError(p)
            if p.suffix.lower()=='.png':
                # Trim only empty viewport margins; the original evidence stays intact.
                im=Image.open(p).convert('RGB');bg='white' if 'browser' in p.stem else im.getpixel((im.width-1,im.height-1))
                mask=ImageChops.difference(im,Image.new('RGB',im.size,bg)).convert('L').point(lambda v:255 if v>20 else 0)
                box=mask.getbbox()
                if box:
                    box=(max(0,box[0]-10),max(0,box[1]-10),min(im.width,box[2]+20),min(im.height,box[3]+20))
                    crop=ROOT/'tmp/report-inputs'/p.name;crop.parent.mkdir(parents=True,exist_ok=True)
                    im.crop(box).save(crop);p=crop
            pp=para(d);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True
            pp.add_run().add_picture(str(p),width=Cm(16.2))
            cp=para(d,item['caption'],'Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading('Ответы на контрольные вопросы',level=1)
    for i,(title,answer) in enumerate(QUESTIONS[str(n)],1):
        d.add_heading(f'{i} {title}',level=2);para(d,answer)
    d.add_heading('Выводы',level=1)
    for text in spec['conclusions']:para(d,text)
    d.add_heading('Репозиторий и источники',level=1)
    for i,text in enumerate(spec['sources']):
        p=para(d,text);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
        p.paragraph_format.keep_with_next=i<len(spec['sources'])-1
        for run in p.runs:run.font.size=Pt(11)
    for appendix in spec.get('appendices',[]):
        d.add_page_break();d.add_heading(appendix['title'],level=1)
        path=ROOT/appendix['file'];body=path.read_text(encoding='utf-8')
        if 'start' in appendix:body=body.split(appendix['start'],1)[-1]
        code(d,body)
    d.core_properties.author='Кудояров Денис Рустемович'
    d.core_properties.title=f'Практическая работа {n} {spec["title"]}'
    path=OUT/f'Практическая работа {n}.docx';d.save(path);print(path)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    specs=json.loads((ROOT/'lab_tools/report_specs.json').read_text(encoding='utf-8'))
    for n,spec in specs.items():create(int(n),spec)
