"""Typeset the existing lab evidence using applicable GOST 7.32-2017 rules."""
from pathlib import Path
import json,re,sys
from PIL import Image,ImageChops
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/reports'
EVID=ROOT/'reports/practices-5-8/evidence'
GOST_URL='https://protect.gost.ru/gost/details/7d280e43-7036-4a69-8e6e-d15867028343'
ACCESS_DATE='23.10.2026'
WIDTH=16.5

def font(style,name='Times New Roman',size=14,bold=False):
    style.font.name=name;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    props=style.element.get_or_add_rPr();fonts=props.rFonts
    for key in list(fonts.attrib):
        if 'theme' in key.lower():del fonts.attrib[key]
    for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),name)
    for key in ('color','u'):
        for el in props.findall(qn('w:'+key)):
            for attr in list(el.attrib):
                if 'theme' in attr.lower():del el.attrib[attr]
    ppr=style.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)

def para(doc,text='',style=None,indent=True,align=None):
    p=doc.add_paragraph(text,style)
    p.paragraph_format.widow_control=True
    if not indent:p.paragraph_format.first_line_indent=Cm(0)
    if align is not None:p.alignment=align
    return p

def heading(doc,text,level=1):
    p=doc.add_heading(text,level)
    p.paragraph_format.widow_control=True
    return p

def structural(doc,text,style='GOST Structural'):
    return para(doc,text,style,False,WD_ALIGN_PARAGRAPH.CENTER)

def listing(doc,text,evidence=False):
    lines=text.replace('\r\n','\n').rstrip().splitlines()
    if evidence:
        lines=[line for line in lines if line.strip() and not re.match(r'^\s*(% Total|Dload|\d+\s+\d+\s+\d+\s+\d+)',line)]
    for i,line in enumerate(lines):
        p=para(doc,line.expandtabs(4),'Code',False,WD_ALIGN_PARAGRAPH.LEFT)
        # Small command groups stay readable together; long listings paginate normally.
        p.paragraph_format.keep_with_next=len(lines)<=8 and i<len(lines)-1

def config(doc):
    s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2)
    s.header_distance=Cm(1);s.footer_distance=Cm(1)
    normal=doc.styles['Normal'];font(normal)
    pf=normal.paragraph_format;pf.first_line_indent=Cm(1.25);pf.line_spacing=1.5
    pf.space_before=Pt(0);pf.space_after=Pt(0);pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.widow_control=True
    for name in ('Title','Subtitle','Heading 1','Heading 2','Caption'):
        font(doc.styles[name])
    for name in ('Title','Heading 1','Heading 2'):doc.styles[name].font.bold=True
    for name in ('Title','Subtitle'):
        p=doc.styles[name].paragraph_format;p.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.space_before=Pt(0);p.space_after=Pt(6);p.line_spacing=1.5;p.keep_with_next=True
    for name in ('Heading 1','Heading 2'):
        p=doc.styles[name].paragraph_format;p.first_line_indent=Cm(1.25);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.space_before=Pt(12 if name=='Heading 2' else 0);p.space_after=Pt(12)
        p.line_spacing=1.5;p.keep_with_next=True;p.keep_together=True
        p.page_break_before=name=='Heading 1'
        lang=OxmlElement('w:suppressAutoHyphens');doc.styles[name].element.get_or_add_pPr().append(lang)
    for name in ('GOST Structural','GOST TOC Title','GOST Appendix'):
        st=doc.styles.add_style(name,1);st.base_style=normal;font(st,bold=True)
        p=st.paragraph_format;p.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.space_before=Pt(0);p.space_after=Pt(18);p.keep_with_next=True;p.keep_together=True;p.page_break_before=True
        if name!='GOST TOC Title':
            el=OxmlElement('w:outlineLvl');el.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(el)
    st=doc.styles.add_style('Code',1);st.base_style=normal;font(st,'Consolas',12)
    p=st.paragraph_format;p.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.line_spacing=1.5
    p.space_before=Pt(0);p.space_after=Pt(0);p.keep_together=True
    font(doc.styles['Caption'],size=12)
    p=doc.styles['Caption'].paragraph_format;p.first_line_indent=Cm(0);p.line_spacing=1;p.space_before=Pt(6);p.space_after=Pt(12)
    for i in (1,2):
        name=f'TOC {i}'
        st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
        st.base_style=normal;font(st,size=12)
        p=st.paragraph_format;p.first_line_indent=Cm(0);p.left_indent=Cm(.5 if i==2 else 0)
        p.line_spacing=1;p.space_before=Pt(0);p.space_after=Pt(5);p.keep_with_next=False
        p.tab_stops.add_tab_stop(Cm(WIDTH),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    if 'Hyperlink' in doc.styles:
        font(doc.styles['Hyperlink']);doc.styles['Hyperlink'].font.underline=False
    s.different_first_page_header_footer=True
    p=s.first_page_footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    run=p.add_run('Москва 2026');run.font.name='Times New Roman';run.font.size=Pt(14)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.line_spacing=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)

class Report:
    def __init__(self,n,spec):
        self.n=n;self.spec=spec;self.doc=Document();self.figure=0;self.table_number=0
        config(self.doc)

    def table(self,headers,rows,title,widths=None):
        d=self.doc;self.table_number+=1;num=self.table_number
        para(d,f'Данные приведены в таблице {num}.').paragraph_format.keep_with_next=True
        p=para(d,f'Таблица {num} — {title}',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
        if widths:
            factor=WIDTH/sum(widths);widths=[w*factor for w in widths]
        elif len(headers)==2:widths=[5,WIDTH-5]
        elif len(headers)==3:widths=[5,1.5,WIDTH-6.5]
        elif len(headers)==4:widths=[2,2.5,1.5,WIDTH-6]
        else:widths=[WIDTH/len(headers)]*len(headers)
        t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for cell,value in zip(t.rows[0].cells,headers):cell.text=str(value)
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for ri,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'8');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or (ci>0 and len(str(cell.text))<18) else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(12);run.font.bold=ri==0
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        para(d,'',indent=False).paragraph_format.space_after=Pt(3)

    def image(self,item):
        d=self.doc;self.figure+=1;num=self.figure;source=EVID/item['file']
        im=Image.open(source).convert('RGB');bg='white' if 'browser' in source.stem else im.getpixel((im.width-1,im.height-1))
        box=ImageChops.difference(im,Image.new('RGB',im.size,bg)).convert('L').point(lambda v:255 if v>20 else 0).getbbox()
        if box:
            box=(max(0,box[0]-10),max(0,box[1]-10),min(im.width,box[2]+20),min(im.height,box[3]+20));im=im.crop(box)
        crop=ROOT/'tmp/report-inputs'/source.name;crop.parent.mkdir(parents=True,exist_ok=True);im.save(crop)
        title=re.sub(r'^Рисунок\s+\d+\s*[—.-]?\s*','',item['caption']).strip()
        para(d,f'Результат проверки представлен на рисунке {num}.')
        p=para(d,'',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
        width=min(16,18*im.width/im.height);p.add_run().add_picture(str(crop),width=Cm(width))
        para(d,f'Рисунок {num} — {title}','Caption',False,WD_ALIGN_PARAGRAPH.CENTER)

    def evidence(self,item):
        text=(EVID/item['file']).read_text(encoding='utf-8',errors='replace')
        if item['file']=='snapshots.txt':
            rows=re.findall(r'Name: (.*?) \(UUID: ([^)]+)\)',text)
            self.table(['Снимок','UUID'],rows,'Контрольные точки виртуальной машины',[6,10.5])
            return
        if 'start' in item:
            if item['start'] not in text:raise ValueError(item['start'])
            text=text.split(item['start'],1)[1].lstrip('\n')
        if 'end' in item:
            if item['end'] not in text:raise ValueError(item['end'])
            text=text.split(item['end'],1)[0].rstrip()
        if 'lines' in item:text='\n'.join(text.splitlines()[:item['lines']])
        p=para(self.doc,item['caption']);p.paragraph_format.keep_with_next=True
        listing(self.doc,text,True)

    def sources(self):
        n=self.n
        result=[
            f'Практическая работа № {n} : методические указания по дисциплине «Инструменты девопс» / Д. М. Гиматдинов. — 2026. — { {5:18,6:20,7:14}[n] } с. — Текст : электронный.',
            f'Tools_DevOps : учебный репозиторий. — URL: https://github.com/Hehekutov/Tools_DevOps (дата обращения: {ACCESS_DATE}). — Текст : электронный.',
            f'Практические работы № 5–7 : Pull Request № 3. — URL: https://github.com/Hehekutov/Tools_DevOps/pull/3 (дата обращения: {ACCESS_DATE}). — Текст : электронный.',
            f'ГОСТ 7.32–2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления. — URL: {GOST_URL} (дата обращения: {ACCESS_DATE}). — Текст : электронный.'
        ]
        for source in self.spec['sources']:
            m=re.search(r'https?://\S+',source)
            if not m or 'Tools_DevOps' in m[0] or 'protect.gost.ru' in m[0]:continue
            url=m[0].rstrip('.,');title=source[:m.start()].strip().rstrip('.')
            result.append(f'{title}. — URL: {url} (дата обращения: {ACCESS_DATE}). — Текст : электронный.')
        return result

    def build(self):
        d=self.doc;s=self.spec;n=self.n;sources=self.sources()
        for text in ('МИРЭА — Российский технологический университет','Институт перспективных технологий и индустриального программирования','Кафедра индустриального программирования'):
            p=para(d,text,indent=False,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.line_spacing=1
        for _ in range(3):para(d,'',indent=False)
        para(d,f'ОТЧЕТ\nПО ПРАКТИЧЕСКОЙ РАБОТЕ № {n}','Title',False,WD_ALIGN_PARAGRAPH.CENTER)
        para(d,s['title'],'Subtitle',False,WD_ALIGN_PARAGRAPH.CENTER)
        para(d,'по дисциплине «Инструменты девопс»',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(3):para(d,'',indent=False)
        for text in ('Выполнил студент группы ЭФБО-18-24','Кудояров Денис Рустемович'):
            para(d,text,indent=False,align=WD_ALIGN_PARAGRAPH.RIGHT)
        structural(d,'СОДЕРЖАНИЕ','GOST TOC Title')
        p=para(d,'',indent=False);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');field.set(qn('w:dirty'),'true');p._p.append(field)
        structural(d,'ВВЕДЕНИЕ')
        para(d,'Практическая работа выполнена по методическим указаниям [1]. Исходные тексты и материалы проверок размещены в репозитории [2], итоговые изменения представлены в Pull Request [3].')
        for text in s['intro']:para(d,text)
        para(d,'При оформлении применены правила ГОСТ 7.32–2017 для текста, нумерации, иллюстраций, таблиц и приложений [4] с сохранением структуры учебной практической работы.')
        if len(sources)>4:para(d,f'При настройке и проверке программных средств использована документация, приведённая в источниках [5–{len(sources)}].')
        heading(d,'1 Рабочее окружение')
        self.table(['Параметр','Значение'],s['environment'],'Параметры рабочего окружения',[5,11.5])
        heading(d,'2 Выполнение практической работы')
        for i,section in enumerate(s['sections'],1):
            title=re.sub(r'^\d+\s+','',section['title']);heading(d,f'2.{i} {title}',2)
            for text in section.get('paragraphs',[]):para(d,text)
            if 'table' in section:
                t=section['table'];self.table(t['headers'],t['rows'],title,t.get('widths'))
            if 'commands' in section:listing(d,section['commands'])
            for item in section.get('evidence',[]):self.evidence(item)
            for item in section.get('images',[]):self.image(item)
        heading(d,'3 Ответы на контрольные вопросы')
        questions=json.loads((ROOT/'lab_tools/questions.json').read_text(encoding='utf-8'))
        for i,(question,answer) in enumerate(questions[str(n)],1):
            heading(d,f'3.{i} {question}',2);para(d,answer)
        structural(d,'ЗАКЛЮЧЕНИЕ')
        for text in s['conclusions']:para(d,text)
        if s.get('appendices'):
            letters='АБВГДЖИКЛМНПРСТУФХЦШЩЭЮЯ'
            refs=', '.join(letters[i] for i in range(len(s['appendices'])))
            para(d,f'Полные тексты сценариев, конфигураций и дополнительные материалы приведены в приложениях {refs}.')
        structural(d,'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ')
        for i,text in enumerate(sources,1):para(d,f'{i}. {text}')
        for i,appendix in enumerate(s.get('appendices',[])):
            label='АБВГДЖИКЛМНПРСТУФХЦШЩЭЮЯ'[i]
            title=re.sub(r'^Приложение\s+[А-ЯA-Z][.]?\s*','',appendix['title'])
            structural(d,f'ПРИЛОЖЕНИЕ {label}\n{title}','GOST Appendix')
            text=(ROOT/appendix['file']).read_text(encoding='utf-8')
            if 'start' in appendix:text=text.split(appendix['start'],1)[-1].lstrip()
            listing(d,text)
        d.core_properties.author='Кудояров Денис Рустемович'
        d.core_properties.title=f'Практическая работа № {n}. {s["title"]}'
        d.core_properties.subject='Учебный отчёт; оформление по правилам ГОСТ 7.32-2017'
        path=OUT/f'Практическая работа {n}.docx';d.save(path);print(path,flush=True)

def build_all():
    sys.stdout.reconfigure(encoding='utf-8');OUT.mkdir(parents=True,exist_ok=True)
    specs=json.loads((ROOT/'lab_tools/report_specs.json').read_text(encoding='utf-8'))
    for n,spec in specs.items():Report(int(n),spec).build()

if __name__=='__main__':build_all()
